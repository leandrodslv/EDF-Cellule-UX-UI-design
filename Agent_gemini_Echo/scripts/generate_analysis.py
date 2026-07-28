"""
ECHO — Générateur d'analyse Excel UXR
Produit un classeur complet à partir des réponses exportées depuis Microsoft Forms.

Usage:
    python generate_analysis.py <responses.xlsx> <project_name>

Exemple:
    python generate_analysis.py onboarding-edf-responses.xlsx "Onboarding EDF"
"""

import sys
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.utils import get_column_letter
from datetime import datetime


# ─── PALETTE ECHO ─────────────────────────────────────────────────────────────

C = {
    "navy":        "1E3A5F",
    "blue":        "2563EB",
    "blue_light":  "EFF6FF",
    "teal":        "0D9488",
    "teal_light":  "F0FDFA",
    "amber":       "D97706",
    "amber_light": "FFFBEB",
    "red":         "DC2626",
    "red_light":   "FEF2F2",
    "gray":        "6B7280",
    "gray_light":  "F9FAFB",
    "white":       "FFFFFF",
    "border":      "E5E7EB",
}

FONT_NAME = "Arial"


# ─── HELPERS ──────────────────────────────────────────────────────────────────

def solid(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def f(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT_NAME, bold=bold, size=size, color=color, italic=italic)

def al(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def thin_border():
    s = Side(style="thin", color=C["border"])
    return Border(top=s, left=s, right=s, bottom=s)

def sc(ws, row, col, value, bold=False, size=10, color="000000",
       bg=None, h="left", wrap=False, italic=False, fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = f(bold=bold, size=size, color=color, italic=italic)
    cell.alignment = al(h, "center", wrap)
    if bg:
        cell.fill = solid(bg)
    if fmt:
        cell.number_format = fmt
    cell.border = thin_border()
    return cell

def header_row(ws, row, labels, bg=None, height=26):
    bg = bg or C["blue"]
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = f(bold=True, size=10, color="FFFFFF")
        c.fill = solid(bg)
        c.alignment = al("center", "center", wrap=True)
        c.border = thin_border()
    ws.row_dimensions[row].height = height

def section_title(ws, row, col, title, span=1, bg=None):
    bg = bg or C["navy"]
    ws.merge_cells(start_row=row, start_column=col,
                   end_row=row, end_column=col + span - 1)
    c = ws.cell(row=row, column=col, value=title)
    c.font = f(bold=True, size=11, color="FFFFFF")
    c.fill = solid(bg)
    c.alignment = al("left", "center")
    ws.row_dimensions[row].height = 26
    return c


# ─── LECTURE DONNÉES ──────────────────────────────────────────────────────────

def load_responses(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value
               for c in range(1, ws.max_column + 1)
               if ws.cell(1, c).value]
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in row):
            data.append(list(row[:len(headers)]))
    return headers, data


# ─── DÉTECTION TYPES ──────────────────────────────────────────────────────────

LIKERT_MAP = {
    "Pas du tout d'accord": 1, "Pas d'accord": 2, "Neutre": 3,
    "D'accord": 4, "Tout à fait d'accord": 5,
    "Jamais": 1, "Rarement": 2, "Parfois": 3, "Souvent": 4, "Toujours": 5,
    "Pas du tout important": 1, "Peu important": 2,
    "Important": 4, "Très important": 5,
    "Très insatisfait": 1, "Insatisfait": 2,
    "Satisfait": 4, "Très satisfait": 5,
    "1": 1, "2": 2, "3": 3, "4": 4, "5": 5,
}

def to_num(val):
    if val is None:
        return None
    if str(val) in LIKERT_MAP:
        return LIKERT_MAP[str(val)]
    try:
        return float(val)
    except (TypeError, ValueError):
        return None

def detect_type(col_idx, data):
    vals = [r[col_idx] for r in data if r[col_idx] is not None]
    if not vals:
        return "empty"
    if any(str(v) in LIKERT_MAP for v in vals):
        return "likert"
    unique = list(set(str(v) for v in vals))
    if len(unique) <= 8:
        return "qcm"
    if all(isinstance(v, (int, float)) for v in vals):
        return "numeric"
    return "text"


# ─── ONGLET DONNÉES BRUTES ────────────────────────────────────────────────────

def sheet_raw(wb, headers, data, project_name):
    ws = wb.create_sheet("📋 Données brutes")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    n_cols = len(headers)
    ws.merge_cells(f"A1:{get_column_letter(n_cols)}1")
    t = ws["A1"]
    t.value = f"DONNÉES BRUTES — {project_name.upper()}"
    t.font = f(bold=True, size=12, color="FFFFFF")
    t.fill = solid(C["navy"])
    t.alignment = al("center", "center")
    ws.row_dimensions[1].height = 30

    header_row(ws, 2, headers, bg=C["blue"])
    for col in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24

    for r_idx, row_data in enumerate(data, 3):
        bg = C["white"] if r_idx % 2 == 0 else C["gray_light"]
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = solid(bg)
            cell.font = f(size=9)
            cell.alignment = al("left", "center", wrap=True)
            cell.border = thin_border()
        ws.row_dimensions[r_idx].height = 18


# ─── ONGLET PROFILS ───────────────────────────────────────────────────────────

def sheet_profiles(wb, headers, data, project_name):
    ws = wb.create_sheet("👤 Profils")
    ws.sheet_view.showGridLines = False

    qcm_cols = [(i, h) for i, h in enumerate(headers)
                if detect_type(i, data) == "qcm"]

    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = f"PROFILS DES RÉPONDANTS — {project_name.upper()}"
    t.font = f(bold=True, size=12, color="FFFFFF")
    t.fill = solid(C["navy"])
    t.alignment = al("center")
    ws.row_dimensions[1].height = 30

    row = 3
    n = len(data)
    sc(ws, row, 1, f"Total répondants : {n}", bold=True, size=11,
       color=C["navy"], bg=C["blue_light"])
    ws.merge_cells(f"A{row}:D{row}")
    row += 2

    for col_idx, question in qcm_cols:
        section_title(ws, row, 1, question, span=5)
        row += 1
        header_row(ws, row, ["Réponse", "N", "%", "Distribution"], bg=C["teal"])
        row += 1

        vals = [r[col_idx] for r in data if r[col_idx] is not None]
        counts = {}
        for v in vals:
            counts[str(v)] = counts.get(str(v), 0) + 1

        chart_start = row
        for option, count in sorted(counts.items(),
                                    key=lambda x: x[1], reverse=True):
            pct = count / n if n > 0 else 0
            bg = C["white"] if row % 2 == 0 else C["gray_light"]
            sc(ws, row, 1, option, bg=bg)
            sc(ws, row, 2, count, h="center", bg=bg)
            sc(ws, row, 3, pct, h="center", bg=bg, fmt="0.0%")
            # Barre visuelle textuelle
            bar_len = int(pct * 20)
            sc(ws, row, 4, "█" * bar_len, bg=bg, color=C["teal"])
            ws.row_dimensions[row].height = 20
            row += 1

        chart_end = row - 1

        # Graphique camembert
        if chart_end >= chart_start + 1:
            pie = PieChart()
            pie.title = str(question)[:40]
            pie.style = 10
            pie.width = 12
            pie.height = 8
            labels = Reference(ws, min_col=1, min_row=chart_start,
                                max_row=chart_end)
            data_ref = Reference(ws, min_col=2, min_row=chart_start,
                                 max_row=chart_end)
            pie.add_data(data_ref)
            pie.set_categories(labels)
            ws.add_chart(pie, f"F{chart_start}")

        row += 2

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 22


# ─── ONGLET SATISFACTION ──────────────────────────────────────────────────────

def sheet_satisfaction(wb, headers, data, project_name):
    ws = wb.create_sheet("⭐ Satisfaction")
    ws.sheet_view.showGridLines = False

    likert_cols = [(i, h) for i, h in enumerate(headers)
                   if detect_type(i, data) == "likert"]

    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value = f"ANALYSE SATISFACTION / LIKERT — {project_name.upper()}"
    t.font = f(bold=True, size=12, color="FFFFFF")
    t.fill = solid(C["navy"])
    t.alignment = al("center")
    ws.row_dimensions[1].height = 30

    row = 3

    if not likert_cols:
        ws.cell(row, 1, "Aucune question Likert détectée.")
        return

    section_title(ws, row, 1, "SCORES PAR QUESTION", span=8)
    row += 1
    header_row(ws, row,
        ["Question", "Moy. /5", "Médiane", "Écart-type",
         "% ≥ 4 (positif)", "% ≤ 2 (négatif)", "Polarisation", "Signal"],
        bg=C["blue"])
    row += 1

    scores = []
    chart_labels = []
    chart_values = []
    chart_row_start = row

    for col_idx, question in likert_cols:
        nums = [to_num(r[col_idx]) for r in data if r[col_idx] is not None]
        nums = [n for n in nums if n is not None]
        if not nums:
            continue

        n = len(nums)
        avg = sum(nums) / n
        nums_s = sorted(nums)
        med = nums_s[n // 2] if n % 2 else (nums_s[n//2-1] + nums_s[n//2]) / 2
        std = (sum((x - avg)**2 for x in nums) / n) ** 0.5
        pct_pos = len([x for x in nums if x >= 4]) / n
        pct_neg = len([x for x in nums if x <= 2]) / n
        polar = "⚠️ Oui" if std > 1.2 else "Non"
        signal = ("🟢 Positif" if avg >= 4 else
                  "🟡 Neutre" if avg >= 3 else "🔴 Négatif")

        bg = C["white"] if row % 2 == 0 else C["gray_light"]
        sc(ws, row, 1, str(question)[:60], bg=bg, wrap=True)
        sc(ws, row, 2, round(avg, 2), h="center", bg=bg, fmt="0.00")
        sc(ws, row, 3, round(med, 1), h="center", bg=bg)
        sc(ws, row, 4, round(std, 2), h="center", bg=bg)
        sc(ws, row, 5, pct_pos, h="center", bg=bg, fmt="0%")
        sc(ws, row, 6, pct_neg, h="center", bg=bg, fmt="0%")
        pol_c = ws.cell(row=row, column=7, value=polar)
        pol_c.font = f(color=C["amber"] if "Oui" in polar else C["teal"])
        pol_c.fill = solid(bg)
        pol_c.alignment = al("center")
        pol_c.border = thin_border()
        sig_c = ws.cell(row=row, column=8, value=signal)
        sig_c.font = f(color=(C["teal"] if "Positif" in signal else
                               C["amber"] if "Neutre" in signal else C["red"]))
        sig_c.fill = solid(bg)
        sig_c.alignment = al("center")
        sig_c.border = thin_border()

        ws.row_dimensions[row].height = 36
        scores.append(round(avg, 2))
        chart_labels.append(str(question)[:30])
        chart_values.append(round(avg, 2))
        row += 1

    # Score global
    row += 1
    section_title(ws, row, 1, "SCORE GLOBAL", span=3, bg=C["teal"])
    row += 1
    if scores:
        g = sum(scores) / len(scores)
        sc(ws, row, 1, "Score moyen global", bold=True, color=C["navy"], bg=C["teal_light"])
        score_c = ws.cell(row=row, column=2, value=round(g, 2))
        score_c.font = f(bold=True, size=14,
            color=(C["teal"] if g >= 4 else C["amber"] if g >= 3 else C["red"]))
        score_c.fill = solid(C["teal_light"])
        score_c.alignment = al("center")
        score_c.border = thin_border()
        interp = ("Très satisfaisant ✅" if g >= 4.5 else
                  "Satisfaisant 🟢" if g >= 4 else
                  "Neutre 🟡" if g >= 3 else "Insatisfaisant 🔴")
        sc(ws, row, 3, f"/ 5.0 — {interp}", italic=True, color=C["gray"],
           bg=C["teal_light"])
        row += 2

        # Données graphique barre
        data_col = 10
        ws.cell(row, data_col, "Question").font = f(bold=True, size=9)
        ws.cell(row, data_col + 1, "Score moyen").font = f(bold=True, size=9)
        row += 1
        for label, val in zip(chart_labels, chart_values):
            ws.cell(row, data_col, label)
            ws.cell(row, data_col + 1, val)
            row += 1

        bar = BarChart()
        bar.type = "bar"
        bar.style = 10
        bar.title = "Scores moyens Likert"
        bar.y_axis.title = "Score / 5"
        bar.width = 20
        bar.height = 14
        data_ref = Reference(ws, min_col=data_col + 1,
                              min_row=row - len(chart_values) - 1,
                              max_row=row - 1)
        cats = Reference(ws, min_col=data_col,
                         min_row=row - len(chart_values),
                         max_row=row - 1)
        bar.add_data(data_ref, titles_from_data=True)
        bar.set_categories(cats)
        ws.add_chart(bar, f"A{row + 1}")

    ws.column_dimensions["A"].width = 45
    for col in ["B","C","D","E","F","G","H"]:
        ws.column_dimensions[col].width = 16


# ─── ONGLET PROBLÈMES ─────────────────────────────────────────────────────────

def sheet_problems(wb, headers, data, project_name):
    ws = wb.create_sheet("⚡ Problèmes")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:H1")
    t = ws["A1"]
    t.value = f"MATRICE IMPACT × FRÉQUENCE — {project_name.upper()}"
    t.font = f(bold=True, size=12, color="FFFFFF")
    t.fill = solid(C["navy"])
    t.alignment = al("center")
    ws.row_dimensions[1].height = 30

    row = 3
    note = ws.cell(row, 1,
        "💡 Complète les colonnes Fréquence et Impact (1=faible → 5=fort). "
        "La priorité se calcule automatiquement.")
    note.font = f(size=9, italic=True, color=C["gray"])
    ws.merge_cells(f"A{row}:H{row}")
    row += 2

    section_title(ws, row, 1, "PROBLÈMES IDENTIFIÉS", span=7)
    row += 1
    header_row(ws, row,
        ["Problème", "Source", "Fréq.\n(1-5)", "Impact\n(1-5)",
         "Score", "Priorité", "Action recommandée"],
        bg=C["amber"], height=36)
    row += 1

    examples = [
        ("Processus trop long", "Entretien + Questionnaire", 4, 4),
        ("Manque de feedback visuel", "Entretien", 3, 4),
        ("Navigation peu claire sur mobile", "Questionnaire", 4, 3),
        ("", "", None, None), ("", "", None, None),
        ("", "", None, None), ("", "", None, None),
        ("", "", None, None), ("", "", None, None),
    ]

    for i, (prob, source, freq, impact) in enumerate(examples):
        r = row + i
        bg = C["white"] if i % 2 == 0 else C["gray_light"]
        sc(ws, r, 1, prob, bg=bg)
        sc(ws, r, 2, source, bg=bg)
        # Colonnes à remplir manuellement (fond ambre)
        for col in [3, 4]:
            val = freq if col == 3 else impact
            c = ws.cell(row=r, column=col, value=val)
            c.fill = solid(C["amber_light"])
            c.alignment = al("center")
            c.border = thin_border()
        # Score auto
        score_c = ws.cell(row=r, column=5,
            value=f'=IF(AND(C{r}<>"",D{r}<>""),C{r}*D{r},"")')
        score_c.fill = solid(bg)
        score_c.alignment = al("center")
        score_c.font = f(bold=True)
        score_c.border = thin_border()
        # Priorité auto
        prio_c = ws.cell(row=r, column=6,
            value=(f'=IF(E{r}="","",IF(E{r}>=16,"🔴 P1 Critique",'
                   f'IF(E{r}>=9,"🟠 P2 Important",'
                   f'IF(E{r}>=4,"🟡 P3 Moyen","⚪ P4 Faible"))))'))
        prio_c.fill = solid(bg)
        prio_c.alignment = al("center")
        prio_c.border = thin_border()
        sc(ws, r, 7, "", bg=bg)
        ws.row_dimensions[r].height = 22

    row += len(examples) + 2

    # Légende
    section_title(ws, row, 1, "LÉGENDE PRIORITÉ", span=3, bg=C["gray"])
    row += 1
    for label, desc, color in [
        ("🔴 P1 Critique", "Score ≥ 16 — traiter dans le sprint en cours", C["red"]),
        ("🟠 P2 Important", "Score ≥ 9 — planifier dans le prochain sprint", C["amber"]),
        ("🟡 P3 Moyen", "Score ≥ 4 — backlog prioritaire", C["gray"]),
        ("⚪ P4 Faible", "Score < 4 — backlog", C["gray"]),
    ]:
        sc(ws, row, 1, label, bold=True, color=color, bg=C["gray_light"])
        sc(ws, row, 2, desc, italic=True, color=C["gray"], bg=C["gray_light"])
        ws.merge_cells(f"B{row}:C{row}")
        row += 1

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 38


# ─── ONGLET CORRÉLATIONS ──────────────────────────────────────────────────────

def sheet_correlations(wb, headers, data, project_name):
    ws = wb.create_sheet("🔗 Corrélations")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    t = ws["A1"]
    t.value = f"CORRÉLATIONS & CROISEMENTS — {project_name.upper()}"
    t.font = f(bold=True, size=12, color="FFFFFF")
    t.fill = solid(C["navy"])
    t.alignment = al("center")
    ws.row_dimensions[1].height = 30

    row = 3

    likert_cols = [(i, h) for i, h in enumerate(headers)
                   if detect_type(i, data) == "likert"]
    qcm_cols = [(i, h) for i, h in enumerate(headers)
                if detect_type(i, data) == "qcm"]

    if not (likert_cols and qcm_cols):
        ws.cell(row, 1,
            "Pas assez de variables pour calculer des corrélations. "
            "Il faut au moins 1 colonne Likert et 1 QCM.")
        return

    section_title(ws, row, 1, "SCORES MOYENS PAR SEGMENT", span=5)
    row += 1
    note = ws.cell(row, 1,
        "Comment lire : chaque tableau montre comment le score varie selon le profil du répondant.")
    note.font = f(size=9, italic=True, color=C["gray"])
    ws.merge_cells(f"A{row}:E{row}")
    row += 2

    combos = [(q, l) for q in qcm_cols[:2] for l in likert_cols[:2]]

    for (qcm_i, qcm_q), (lik_i, lik_q) in combos:
        title = f"{str(qcm_q)[:30]} → {str(lik_q)[:30]}"
        section_title(ws, row, 1, title, span=4, bg=C["teal"])
        row += 1
        header_row(ws, row,
            ["Segment", "N", "Score moyen /5", "Écart vs global"],
            bg=C["blue"])
        row += 1

        segments = {}
        all_nums = []
        for r_data in data:
            seg = r_data[qcm_i]
            num = to_num(r_data[lik_i])
            if seg is not None and num is not None:
                segments.setdefault(str(seg), []).append(num)
                all_nums.append(num)

        g_avg = sum(all_nums) / len(all_nums) if all_nums else 0

        for seg, nums in sorted(segments.items(),
                                 key=lambda x: sum(x[1])/len(x[1]), reverse=True):
            avg = sum(nums) / len(nums)
            delta = avg - g_avg
            bg = C["white"] if row % 2 == 0 else C["gray_light"]
            sc(ws, row, 1, seg, bg=bg)
            sc(ws, row, 2, len(nums), h="center", bg=bg)
            avg_c = ws.cell(row=row, column=3, value=round(avg, 2))
            avg_c.fill = solid(bg)
            avg_c.alignment = al("center")
            avg_c.font = f(color=(C["teal"] if avg >= 4 else
                                   C["amber"] if avg >= 3 else C["red"]))
            avg_c.border = thin_border()
            delta_str = ("+" if delta >= 0 else "") + str(round(delta, 2))
            delta_c = ws.cell(row=row, column=4, value=delta_str)
            delta_c.fill = solid(bg)
            delta_c.alignment = al("center")
            delta_c.font = f(color=C["teal"] if delta >= 0 else C["red"])
            delta_c.border = thin_border()
            ws.row_dimensions[row].height = 20
            row += 1

        row += 2

    ws.column_dimensions["A"].width = 35
    for col in ["B","C","D"]:
        ws.column_dimensions[col].width = 16


# ─── ONGLET DASHBOARD ─────────────────────────────────────────────────────────

def sheet_dashboard(wb, headers, data, project_name):
    ws = wb.create_sheet("📊 Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

    n = len(data)

    # Bandeau
    ws.merge_cells("A1:L3")
    t = ws["A1"]
    t.value = f"ECHO · ANALYSE UXR — {project_name.upper()}"
    t.font = f(bold=True, size=18, color="FFFFFF")
    t.fill = solid(C["navy"])
    t.alignment = al("center", "center")
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18

    row = 5
    date_c = ws.cell(row, 1,
        f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')} · {n} répondants")
    date_c.font = f(size=9, italic=True, color=C["gray"])
    ws.merge_cells(f"A{row}:L{row}")
    ws.row_dimensions[row].height = 18
    row += 2

    # KPIs
    section_title(ws, row, 1, "VUE D'ENSEMBLE", span=12, bg=C["blue"])
    row += 1
    ws.row_dimensions[row].height = 18
    ws.row_dimensions[row+1].height = 34

    def kpi_cell(lbl, val, col, bg, text_col):
        c1 = ws.cell(row=row, column=col, value=lbl)
        c1.font = f(size=9, italic=True, color=C["gray"])
        c1.fill = solid(bg)
        c1.alignment = al("center", "center")
        c2 = ws.cell(row=row+1, column=col, value=val)
        c2.font = f(bold=True, size=20, color=text_col)
        c2.fill = solid(bg)
        c2.alignment = al("center", "center")

    kpi_cell("Répondants", n, 1, C["blue_light"], C["navy"])
    n_likert = sum(1 for i, h in enumerate(headers) if detect_type(i, data) == "likert")
    kpi_cell("Questions Likert", n_likert, 3, C["teal_light"], C["teal"])
    n_qcm = sum(1 for i, h in enumerate(headers) if detect_type(i, data) == "qcm")
    kpi_cell("Variables de profil", n_qcm, 5, C["amber_light"], C["amber"])

    row += 3

    # Navigation
    section_title(ws, row, 1, "ONGLETS — NAVIGATION", span=12, bg=C["blue"])
    row += 1

    nav = [
        ("📋 Données brutes", "Export Forms intact — toutes les réponses", C["gray"]),
        ("👤 Profils", "Segmentation des répondants + graphiques camembert", C["blue"]),
        ("⭐ Satisfaction", "Scores Likert, polarisation, score global + graphique barres", C["teal"]),
        ("⚡ Problèmes", "Matrice impact × fréquence — priorité P1 à P4 auto-calculée", C["red"]),
        ("🔗 Corrélations", "Croisements profil × satisfaction — qui est le plus insatisfait ?", C["amber"]),
    ]
    for sheet_name, desc, color in nav:
        bg = C["white"] if row % 2 == 0 else C["gray_light"]
        name_c = ws.cell(row, 1, sheet_name)
        name_c.font = f(bold=True, color=color, size=10)
        name_c.fill = solid(bg)
        name_c.alignment = al("left", "center")
        name_c.border = thin_border()
        ws.merge_cells(f"A{row}:C{row}")
        desc_c = ws.cell(row, 4, desc)
        desc_c.font = f(size=9, italic=True, color=C["gray"])
        desc_c.fill = solid(bg)
        desc_c.alignment = al("left", "center")
        desc_c.border = thin_border()
        ws.merge_cells(f"D{row}:L{row}")
        ws.row_dimensions[row].height = 24
        row += 1

    row += 1

    # Alertes
    section_title(ws, row, 1, "⚠️ SIGNAUX FORTS", span=12, bg=C["amber"])
    row += 1
    for msg in [
        "Les questions marquées ⚠️ dans ⭐ Satisfaction ont un écart-type > 1.2 "
        "→ les répondants sont divisés → priorité pour les entretiens qualitatifs.",
        "Les problèmes P1 dans ⚡ Problèmes (score ≥ 16) sont à traiter "
        "dans le sprint en cours.",
        "Les segments avec un écart négatif dans 🔗 Corrélations "
        "identifient les groupes les plus insatisfaits → cibles prioritaires pour les entretiens.",
    ]:
        c = ws.cell(row, 1, f"→  {msg}")
        c.font = f(size=9, italic=True, color=C["gray"])
        c.alignment = al("left", "center", wrap=True)
        c.fill = solid(C["amber_light"])
        ws.merge_cells(f"A{row}:L{row}")
        ws.row_dimensions[row].height = 32
        row += 1

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 55
    for col in "EFGHIJKL":
        ws.column_dimensions[col].width = 10


# ─── POINT D'ENTRÉE ───────────────────────────────────────────────────────────

def generate_analysis(responses_path, project_name, output_path=None):
    if not output_path:
        slug = project_name.lower().replace(" ", "-")
        output_path = f"{slug}-analyse.xlsx"

    print(f"📂 Lecture : {responses_path}")
    headers, data = load_responses(responses_path)
    print(f"   → {len(headers)} colonnes · {len(data)} répondants")

    wb = Workbook()
    wb.remove(wb.active)

    print("⚙️  Génération des onglets...")
    sheet_raw(wb, headers, data, project_name)
    sheet_profiles(wb, headers, data, project_name)
    sheet_satisfaction(wb, headers, data, project_name)
    sheet_problems(wb, headers, data, project_name)
    sheet_correlations(wb, headers, data, project_name)
    sheet_dashboard(wb, headers, data, project_name)

    wb.save(output_path)
    print(f"\n✅ Classeur généré : {output_path}")
    print(f"   Onglets : {', '.join(ws.title for ws in wb.worksheets)}")
    return output_path


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage : python generate_analysis.py <responses.xlsx> <'Nom du projet'>")
        sys.exit(1)
    generate_analysis(sys.argv[1], sys.argv[2])
